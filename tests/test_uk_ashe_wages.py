"""Synthetic ASHE workbooks only; tests never invent production UK wages."""
import io
import json
import zipfile

from openpyxl import Workbook

from app import uk_ashe_wages as uk


def workbook(value_rows, cv=False):
    wb=Workbook()
    ws=wb.active
    ws.title="All"
    ws.cell(1,1,"Table 14.7b" if cv else "Table 14.7a")
    headers=["Description","Code","(thousand)","Median","change","Mean","change"]
    for c,v in enumerate(headers,1):
        ws.cell(5,c,v)
    r=6
    # >300 four-digit rows are required by the parser.
    codes=[str(i) for i in range(1000,1310)]
    codes.extend(code for code in value_rows if code not in codes)
    for code in codes:
        title=f"Synthetic occupation {code}"
        med=30000
        mean=35000
        if code in value_rows:
            title,med,mean=value_rows[code]
        ws.cell(r,1,title); ws.cell(r,2,code); ws.cell(r,3,10)
        ws.cell(r,4,med); ws.cell(r,5,0); ws.cell(r,6,mean); ws.cell(r,7,0)
        r+=1
    out=io.BytesIO(); wb.save(out); return out.getvalue()


def source_zip(path, *, median=42000, mean=47000, cv_median=4.0, cv_mean=6.0):
    code,title=uk.SOC2020["civil_engineer"]
    fallback_code,fallback_title=uk.SOC2020["accountant"]
    pay={code:(title,median,mean), fallback_code:(fallback_title,40000,45000)}
    cvs={code:(title,cv_median,cv_mean), fallback_code:(fallback_title,3.0,4.0)}
    with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
        z.writestr("PROV Table 14.7a Annual pay - Gross 2025.xlsx",workbook(pay))
        z.writestr("PROV Table 14.7b Annual pay - Gross 2025 CV.xlsx",workbook(cvs,cv=True))


def test_build_snapshot_keeps_direct_soc_and_ons_quality(tmp_path):
    p=tmp_path/"ashe.zip"; source_zip(p)
    obj=uk.build_snapshot(p)
    rows=[r for r in obj["records"] if r["occupation"]=="civil_engineer"]
    assert {r["measure"] for r in rows}=={"median","mean"}
    median=next(r for r in rows if r["measure"]=="median")
    assert median["value"]==42000
    assert median["cv_percent"]==4.0
    assert median["quality"]=="precise"
    assert median["classification"]=="SOC2020:2121"
    assert obj["preferred_measure"]=="median"


def test_cv_over_20_or_suppression_never_becomes_salary(tmp_path):
    p=tmp_path/"ashe.zip"; source_zip(p,cv_median=21.0,mean="x",cv_mean=3.0)
    obj=uk.build_snapshot(p)
    rows=[r for r in obj["records"] if r["occupation"]=="civil_engineer"]
    assert rows==[]
    reasons=[r for r in obj["rejected"] if r["occupation"]=="civil_engineer"]
    assert {r["reason"] for r in reasons}=={"suppressed_or_cv_over_20"}


def test_title_mismatch_fails_mapping_closed(tmp_path):
    p=tmp_path/"ashe.zip"
    code,_=uk.SOC2020["civil_engineer"]
    with zipfile.ZipFile(p,"w",zipfile.ZIP_DEFLATED) as z:
        fallback_code,fallback_title=uk.SOC2020["accountant"]
        z.writestr("Table 14.7a.xlsx",workbook({
            code:("Wrong profession",42000,47000),
            fallback_code:(fallback_title,40000,45000)}))
        z.writestr("Table 14.7b.xlsx",workbook({
            code:("Wrong profession",4,6),
            fallback_code:(fallback_title,3,4)},cv=True))
    obj=uk.build_snapshot(p)
    assert not [r for r in obj["records"] if r["occupation"]=="civil_engineer"]
    assert any(r["occupation"]=="civil_engineer" and r["reason"]=="missing_or_title_mismatch"
               for r in obj["rejected"])


def test_broad_titles_intentionally_have_no_soc2020_mapping():
    for job in ("doctor","nurse","psychologist","teacher","manager",
                "construction_worker","supermarket_worker","industrial_operator"):
        assert job not in uk.SOC2020


def test_snapshot_load_and_read_uses_median_as_preferred(tmp_path,monkeypatch):
    p=tmp_path/"ashe.zip"; source_zip(p)
    obj=uk.build_snapshot(p)
    out=tmp_path/"snapshot.json"; out.write_text(json.dumps(obj),encoding="utf-8")
    db=tmp_path/"wages.sqlite"
    monkeypatch.setattr("app.store.DB_PATH",str(db))
    # uk imported connect directly; patch it too.
    from app import store
    monkeypatch.setattr(uk,"connect",store.connect)
    assert uk.load_snapshot(out)>0
    wage=uk.wages("GB","civil_engineer")
    assert wage["status"]=="available"
    assert wage["preferred_measure"]=="median"
    assert {x["measure"] for x in wage["observations"]}=={"median","mean"}
    unavailable=uk.wages("GB","doctor")
    assert unavailable["status"]=="unavailable"
    assert "one-to-one" in unavailable["reason"]
