"""Read-only audit of Spain INE 2022 Encuesta de Estructura Salarial microdata.

Never export row-level wages, names, IDs, records or sample individuals.
No inference that a 1-/2-/3-digit CNO group equals a detailed occupation.
This auditor DOES NOT IMPORT wages. Inspect the official variable dictionary
before writing source-specific, population-weighted statistical ingestion.

Usage:
    python -m scripts.inspect_es_earnings_2022 --zip EES_2022_OFFICIAL.zip
    python -m scripts.inspect_es_earnings_2022 --csv OFFICIAL.csv
"""
import argparse
import csv
import io
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

ARCHIVE_MAX_BYTES = 550_000_000
MEMBER_MAX_BYTES = 350_000_000
MAX_ROWS = 30_000
MAX_DICTIONARY_ROWS = 1000
MAX_CATEGORIES = 35
MAX_ENTRIES = 500
DATA_SUFFIXES = (".csv", ".txt")
DICTIONARY_SUFFIXES = (".xlsx",)
OCCUPATION = re.compile(r"(cno|ocupac|occup|profesi)", re.I)
SALARY = re.compile(r"(ganan|salari|remun|earning|wage|pay|bruto|annual|anual)", re.I)
WEIGHT = re.compile(r"(factor|ponder|weight|fexp|elevac)", re.I)


def _encoding(raw):
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for enc in ("utf-8", "cp1252"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    raise ValueError("Unknown official CSV character encoding")


def _columns(columns):
    return {
        "occupation_candidates": [c for c in columns if OCCUPATION.search(c)],
        "salary_candidates": [c for c in columns if SALARY.search(c)],
        "weight_candidates": [c for c in columns if WEIGHT.search(c)],
    }


def inspect_csv_file(stream, name, *, max_rows=MAX_ROWS):
    """Read only an upper-bounded head of CSV; never expose individual records."""
    head=stream.read(32_768)
    if not head:
        raise ValueError("Empty microdata CSV")
    encoding=_encoding(head)
    dialect=";"
    probe=head.decode(encoding,errors="replace")
    try:
        dialect=csv.Sniffer().sniff(probe,delimiters=";,\t|").delimiter
    except csv.Error:
        dialect=";" if probe.splitlines()[0].count(";") >= probe.splitlines()[0].count(",") else ","
    # Streaming prefix wrapper: zip readers cannot reliably seek backwards.
    wrapped=io.TextIOWrapper(io.BufferedReader(_PrefixReader(head,stream)),
                             encoding=encoding,errors="replace",newline="")
    reader=csv.DictReader(wrapped,delimiter=dialect)
    fields=reader.fieldnames or []
    if len(fields)<2 or len(set(fields))!=len(fields):
        raise ValueError("CSV header absent/ambiguous; inspect TXT fixed-width dictionary")
    found=_columns(fields)
    codes={c:Counter() for c in found["occupation_candidates"]}
    consumed=0
    for row in reader:
        consumed+=1
        for field,values in codes.items():
            value=(row.get(field) or "").strip()
            if value and len(values)<MAX_CATEGORIES+1:
                values[value]+=1
        if consumed>=max_rows:
            break
    occupation=[]
    for field,values in codes.items():
        observed=list(values)
        looks_four=[v for v in observed if re.fullmatch(r"[0-9]{4}",v)]
        occupation.append({
            "field":field,
            "sample_distinct_values_first_35": observed[:MAX_CATEGORIES],
            "code_lengths":sorted({len(v) for v in observed}),
            "sample_has_four_digit_numeric_codes":bool(looks_four),
            "truncated_categories":len(values)>MAX_CATEGORIES,
        })
    return {
        "file":name,
        "header_columns":fields[:130],
        "header_truncated":len(fields)>130,
        "delimiter":repr(dialect),
        "sample_rows_read":consumed,
        "sample_incomplete":consumed>=max_rows,
        "occupation_fields":occupation,
        "salary_field_candidates":found["salary_candidates"],
        "weight_field_candidates":found["weight_candidates"],
        "provenance_warning":"Column names and sample categories are leads, not approval of CNO-11 granularity or weighting.",
    }


class _PrefixReader(io.RawIOBase):
    """Expose previously sampled bytes followed by the original non-seekable stream."""
    def __init__(self,head,stream):
        self.head=io.BytesIO(head)
        self.stream=stream
    def readable(self):
        return True
    def readinto(self,buffer):
        b=self.head.read(len(buffer))
        if not b:
            b=self.stream.read(len(buffer)) or b""
        buffer[:len(b)]=b
        return len(b)


def inspect_dictionary(stream, name):
    """Only variable names/labels; never dump respondent worksheets."""
    from openpyxl import load_workbook
    workbook=load_workbook(stream,read_only=True,data_only=True)
    try:
        sheets=[]
        for ws in workbook.worksheets[:10]:
            matches=[]
            for index,row in enumerate(ws.iter_rows(values_only=True),1):
                text=" ".join(str(x) for x in row[:12] if x is not None)
                if OCCUPATION.search(text) or SALARY.search(text) or WEIGHT.search(text):
                    matches.append({"row_number":index,
                                    "description":text[:240]})
                if index>=MAX_DICTIONARY_ROWS:
                    break
            sheets.append({"name":ws.title,
                           "matching_dictionary_lines":matches[:70],
                           "matches_truncated":len(matches)>70})
        return {"file":name,"sheets":sheets}
    finally:
        workbook.close()


def inspect_zip(path):
    path=Path(path)
    if not path.is_file() or path.stat().st_size>ARCHIVE_MAX_BYTES:
        raise ValueError("Missing or oversized official EES ZIP")
    reports={"archive":path.name,
             "source":"INE Spain, Encuesta de Estructura Salarial 2022",
             "status":"inspection_only_no_import",
             "files":[],"csv":[],"dictionaries":[],"nested_archives":[],
             "issues":[]}
    with zipfile.ZipFile(path) as archive:
        entries=archive.infolist()
        if len(entries)>MAX_ENTRIES:
            raise ValueError("Too many ZIP entries; audit manually")
        for entry in entries:
            if entry.is_dir():
                continue
            n=entry.filename
            reports["files"].append({"name":n,"size_bytes":entry.file_size})
            if entry.file_size>MEMBER_MAX_BYTES:
                reports["issues"].append("Oversized member skipped: "+n)
                continue
            lower=n.lower()
            if lower.endswith(DATA_SUFFIXES):
                with archive.open(entry) as raw:
                    try:
                        reports["csv"].append(inspect_csv_file(raw,n))
                    except (ValueError,UnicodeError,csv.Error) as exc:
                        reports["issues"].append(n+": "+type(exc).__name__+" (inspect its data dictionary)")
            elif lower.endswith(DICTIONARY_SUFFIXES):
                with archive.open(entry) as raw:
                    try:
                        reports["dictionaries"].append(inspect_dictionary(raw,n))
                    except (ValueError,KeyError,zipfile.BadZipFile) as exc:
                        reports["issues"].append(n+": "+type(exc).__name__)
            elif lower.endswith(".zip"):
                reports["nested_archives"].append(n)
    reports["note"]=(
        "No source wages imported. If published occupation codes have fewer "
        "than four CNO-11 digits, this version cannot establish a specific "
        "doctor/nurse/psychologist wage. Inspect design/weight variables before "
        "aggregating; suppress unsafe small cells, respect licence and cite INE.")
    return reports


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    source=parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--zip",type=Path,help="Original, unmodified official INE EES 2022 ZIP")
    source.add_argument("--csv",type=Path,help="Original microdata CSV for schema-only inspection")
    args=parser.parse_args(argv)
    if args.zip:
        result=inspect_zip(args.zip)
    else:
        with args.csv.open("rb") as file:
            result={"source":"INE EES 2022 (verify official file manually)",
                    "status":"inspection_only_no_import",
                    "csv":[inspect_csv_file(file,args.csv.name)]}
    json.dump(result,sys.stdout,ensure_ascii=False,indent=2)
    print()
    return 0


if __name__=="__main__":
    raise SystemExit(main())
